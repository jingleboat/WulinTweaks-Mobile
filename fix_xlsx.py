# -*- coding: utf-8 -*-
"""修复 Excel：把「答题首选项正确.xlsx」里留空的『选项』列补上（按 mod 本意重排）

原表：TestMod0001/Excel/答题.xlsx（uid / 选项 / 正确选项）
mod 表：NewMyMod001/Excel/答题首选项正确.xlsx
问题：mod 表 225 个数据行里有 215 行『选项』是空的，只写了 RightIndex。
      导出成 patch 后就是 "Selections": []，游戏读表时越界。
"""
import os
import shutil

import openpyxl

BASE = r'F:\SteamLibrary\steamapps\common\WulinSH\Wulin_Data\StreamingAssets\ModEditor\ModEditor\Assets\_Mods'
ORIG = os.path.join(BASE, 'TestMod0001', 'Excel', '答题.xlsx')
MOD = os.path.join(BASE, 'NewMyMod001', 'Excel', '答题首选项正确.xlsx')
BAK = MOD + '.bak'

COL_UID, COL_NAME, COL_SEL, COL_RI = 2, 3, 7, 8   # 1-based


def load_orig():
    ws = openpyxl.load_workbook(ORIG, data_only=True).worksheets[0]
    out = {}
    for r in ws.iter_rows(min_row=4):
        uid = r[COL_UID - 1].value
        sel = r[COL_SEL - 1].value
        ri = r[COL_RI - 1].value
        if uid is None or str(uid) == 'end' or not sel:
            continue
        opts = str(sel).split(',')
        try:
            idx = int(ri)
        except (TypeError, ValueError):
            continue
        out[str(uid)] = (opts, idx)
    return out


def main():
    orig = load_orig()
    print('原表条目:', len(orig))

    if not os.path.exists(BAK):
        shutil.copy2(MOD, BAK)
        print('已备份 ->', BAK)

    wb = openpyxl.load_workbook(MOD)
    ws = wb.worksheets[0]
    filled = 0
    kept = 0
    noorig = []
    for i in range(4, ws.max_row + 1):
        uid = ws.cell(i, COL_UID).value
        if uid is None or str(uid) == 'end':
            continue
        sel = ws.cell(i, COL_SEL).value
        if sel not in (None, ''):
            kept += 1
            continue
        o = orig.get(str(uid))
        if not o:
            noorig.append((i, uid, ws.cell(i, COL_NAME).value))
            continue
        opts, idx = o
        idx = max(0, min(idx, len(opts) - 1))
        newsel = [opts[idx]] + [x for k, x in enumerate(opts) if k != idx]
        ws.cell(i, COL_SEL).value = ','.join(newsel)
        ws.cell(i, COL_RI).value = 0
        filled += 1

    print('补全 %d 行，原本已有选项 %d 行' % (filled, kept))
    if noorig:
        print('原表没有的 uid（保持留空）:', noorig[:10])

    wb.save(MOD)
    print('已写出', MOD, os.path.getsize(MOD), '字节')

    # 复核
    ws2 = openpyxl.load_workbook(MOD, data_only=True).worksheets[0]
    empty = []
    for i in range(4, ws2.max_row + 1):
        uid = ws2.cell(i, COL_UID).value
        if uid is None or str(uid) == 'end':
            continue
        if ws2.cell(i, COL_SEL).value in (None, ''):
            empty.append((i, uid))
    print('复核：仍为空的行 %d %s' % (len(empty), empty[:8]))


if __name__ == '__main__':
    main()
