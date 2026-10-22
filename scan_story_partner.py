# -*- coding: utf-8 -*-
"""解析 ModEditor.zip 里全部 1804 个事件文件，找出所有带「需要伙伴」条件的图，
再与存档打开过的图交叉，回答：哪些剧情会因为「角色在队」被放宽而触发。"""
import zipfile, re, os, sys, io, hashlib, gzip, json
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

ZIP = r"D:\steamgame\steamapps\common\WulinSH\Wulin_Data\StreamingAssets\ModEditor\ModEditor.zip"
ROOT = r"C:\Users\User\AppData\LocalLow\DefaultCompany\Wulin\76561198807754750"
LOC = r"D:\steamgame\steamapps\common\WulinSH\Wulin_Data\StreamingAssets\ModEditor\ModEditor\Assets\GameRes\json\3\LocData.txt"
ST = {0:"NotOpen",1:"Close",2:"Opened",3:"Started",4:"已完成",5:"Failed",6:"TimeOut",7:"CloseAfterOpen"}

def unesc(s):
    if "\\u" in s:
        try: return s.encode().decode("unicode_escape")
        except: return s
    return s

# ---------- 角色名 ----------
ROLE = {}
try:
    from openpyxl import load_workbook
    wb = load_workbook(r"E:\download\xls\1\角色池.xlsx", read_only=True, data_only=True)
    ws = wb.worksheets[0]; rows = list(ws.iter_rows(values_only=True)); wb.close()
    for r in rows[3:]:
        try: u = int(r[1])
        except: continue
        n = r[2] or ""
        if n: ROLE[u] = str(n)
except Exception:
    pass

# ---------- 本地化 ----------
loc = {}
for r in json.loads(open(LOC, "rb").read().decode("utf-8-sig"))["LocData"]:
    un = r.get("uName") or ""
    if un: loc[un] = r.get("SChinese") or ""
def nm(g): return loc.get(g+"_name", "")
def ov(g): return loc.get(g+"_overview1", "") or loc.get(g+"_overview", "")

# ---------- 存档 ----------
def dec(p):
    d = open(p,"rb").read(); iv, ct = d[:16], d[16:]
    k = hashlib.pbkdf2_hmac("sha1", b"Meow", iv, 100, 16)
    c = Cipher(algorithms.AES(k), modes.CBC(iv)).decryptor()
    pt = c.update(ct) + c.finalize(); pt = pt[:-pt[-1]]
    return gzip.decompress(pt) if pt[:2] == b"\x1f\x8b" else pt
def load(slot):
    raw = dec(ROOT+"\\"+slot+"\\SaveObjectStory.save").decode("utf-8")
    raw = re.sub(r'(?<=[{,])\s*(-?\d+)\s*:', lambda m:'"%s":'%m.group(1), raw)
    return json.loads(raw)["Data"]["value"]

saves = {}
for slot in ("Save0","Save2"):
    v = load(slot)
    saves[slot] = {g:d.get("State",0) for g,d in v.get("GraphDataList",{}).items() if isinstance(d,dict)}

# ---------- 解析 zip 里所有 story ----------
z = zipfile.ZipFile(ZIP)
files = [n for n in z.namelist() if "/GameRes/Story/" in n and n.endswith(".asset")]
print("解析事件文件 %d 个…\n" % len(files))

with_partner = {}   # 图名 -> [(partnerID, Negate)]
for n in files:
    try:
        d = z.read(n)
    except Exception:
        continue
    txt = d.decode("utf-8","ignore")
    # 图名 = &11400000 那个 graph 对象的 m_Name
    gm = re.search(r"--- !u!114 &11400000\r?\n(.*?)(?=\r?\n--- |\Z)", txt, re.S)
    if not gm: continue
    mm = re.search(r'm_Name:\s*"?([^"\r\n]*)"?', gm.group(1))
    if not mm: continue
    graph = unesc(mm.group(1).strip('"'))
    blocks = re.split(r"^--- !u!114 &", txt, flags=re.M)[1:]
    hits = []
    for b in blocks:
        mn = re.search(r'm_Name:\s*"?([^"\r\n]*)"?', b)
        if not mn: continue
        if unesc(mn.group(1).strip('"')) != "需要伙伴": continue
        pid = re.search(r'^\s*partnerID:\s*(-?\d+)', b, re.M)
        neg = re.search(r'^\s*Negate:\s*(\d+)', b, re.M)
        if pid:
            hits.append((int(pid.group(1)), int(neg.group(1)) if neg else 0))
    if hits:
        with_partner[graph] = hits

print("含「需要伙伴」条件的事件文件: %d 个\n" % len(with_partner))

# ---------- 交叉存档 ----------
for slot in ("Save0","Save2"):
    gdl = saves[slot]
    print("="*92)
    print("%s：打开过(State!=NotOpen) 且 需要某角色在队 的剧情" % slot)
    print("="*92)
    rows = []
    for g, ps in with_partner.items():
        st = gdl.get(g, 0)
        if st == 0: continue
        names_ = "、".join("%s(%d)%s" % (ROLE.get(p,"?"), p, "取反" if neg else "") for p,neg in ps)
        rows.append((g, st, names_, nm(g), ov(g)))
    rows.sort()
    print("%-20s %-9s %-30s %s" % ("剧情图","状态","需要伙伴条件","名称/描述"))
    print("-"*92)
    for g, st, cond, name, o in rows:
        line = "%-20s %-9s %-30s %s" % (g, ST.get(st), cond, name)
        print(line)
        if o: print("%-62s%s" % ("", o[:64]))
    print("\n合计 %d 个\n" % len(rows))

# ---------- 全量：哪些图要求谁在队（不限存档） ----------
print("="*92)
print("全部含「需要伙伴」的图（不限存档，共 %d）" % len(with_partner))
print("="*92)
bychar = {}
for g, ps in with_partner.items():
    for p,neg in ps:
        bychar.setdefault(p, []).append((g,neg))
for p in sorted(bychar):
    gs = sorted(set(g for g,_ in bychar[p]))
    print("  %-16s (%d) : %s%s" % (ROLE.get(p,"?"), p, ", ".join(gs[:14]), " …" if len(gs)>14 else ""))
