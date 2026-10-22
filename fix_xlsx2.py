# -*- coding: utf-8 -*-
"""把「答题首选项正确.xlsx」里『选项』列的公式替换成静态值

原因：mod 表这一列全是公式，例如
      =F15&"_答案2,"&F15&"_答案1,"&F15&"_答案3,"&F15&"_答案4"
      这些公式从来没有被 Excel 计算过（文件里没有缓存值），
      ModEditor / openpyxl(data_only) 读到的就是 None，
      导出成 patch 后就变成 "Selections": []，游戏答题时 Selections[0] 越界。

做法：按公式里 F{row} + 字面量 的拼接顺序算出结果，写成静态字符串。
"""
import os
import re
import shutil

import openpyxl

BASE = r'F:\SteamLibrary\steamapps\common\WulinSH\Wulin_Data\StreamingAssets\ModEditor\ModEditor\Assets\_Mods'
MOD = os.path.join(BASE, 'NewMyMod001', 'Excel', '答题首选项正确.xlsx')
BAK = MOD + '.bak2'
COL_SEL = 7      # G 列：选项
COL_NAME = 6     # F 列：名称


def eval_formula(f, ws, row):
    if not isinstance(f, str) or not f.startswith('='):
        return f
    body = f[1:]
    tokens = re.findall(r'F\d+|[^&]+', body)
    out = []
    for t in tokens:
        t = t.strip()
        if not t:
            continue
        m = re.fullmatch(r'F(\d+)', t)
        if m:
            v = ws.cell(int(m.group(1)), COL_NAME).value
            out.append('' if v is None else str(v))
        elif t.startswith('"') and t.endswith('"'):
            out.append(t[1:-1])
        else:
            out.append(t)
    return ''.join(out)


def main():
    if not os.path.exists(BAK):
        shutil.copy2(MOD, BAK)
        print('已备份 ->', BAK)

    wb = openpyxl.load_workbook(MOD)      # 保留公式
    ws = wb.worksheets[0]
    done = 0
    samples = []
    for i in range(4, ws.max_row + 1):
        c = ws.cell(i, COL_SEL)
        if isinstance(c.value, str) and c.value.startswith('='):
            val = eval_formula(c.value, ws, i)
            if len(samples) < 3:
                samples.append((i, c.value, val))
            c.value = val
            done += 1
    print('替换公式 %d 个' % done)
    for i, f, v in samples:
        print('  行%d: %s\n     -> %s' % (i, f, v))
    wb.save(MOD)
    print('已写出', MOD, os.path.getsize(MOD), '字节')

    # 复核：data_only 也要能读到
    ws2 = openpyxl.load_workbook(MOD, data_only=True).worksheets[0]
    empty = 0
    for i in range(4, ws2.max_row + 1):
        uid = ws2.cell(i, 2).value
        if uid is None or str(uid) == 'end':
            continue
        if ws2.cell(i, COL_SEL).value in (None, ''):
            empty += 1
    print('复核：uid 行里选项仍为空:', empty)


if __name__ == '__main__':
    main()
