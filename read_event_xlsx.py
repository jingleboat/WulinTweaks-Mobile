# -*- coding: utf-8 -*-
import sys, io
from openpyxl import load_workbook
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

for p in (r"E:\download\xls\1\事件总表.xlsx", r"E:\download\xls\3\本地化数据_事件总表.xlsx"):
    print("=" * 70)
    print(p)
    wb = load_workbook(p, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print(f"--- sheet: {ws.title}  dims={ws.max_row}x{ws.max_column}")
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            cells = ["" if c is None else str(c) for c in row]
            line = " | ".join(cells).rstrip(" |")
            print(f"  {i:3}: {line[:220]}")
            if i > 40:
                print("   ... (truncated)")
                break
    wb.close()
